"""Render a `TailoredResume` into a `.docx` file with python-docx.

The model returns bullets keyed by `accomplishment_id`. This module joins each
one back to its full record in the profile so bullets appear under the employer,
role, and dates they actually belong to — which is what makes the output a
resume rather than a bare list.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from datetime import datetime
from pathlib import Path

# The pip package is `python-docx`, but you import it as `docx`. This local
# module is also named docx.py; because we use absolute imports throughout
# (`from jobtailor.render import docx`), there's no clash.
from docx import Document

from jobtailor.models import Profile, TailoredResume
from jobtailor.render import skills as skills_mod


def _fmt_date(value: str | None) -> str:
    """Turn "2024-06" into "Jun 2024"; pass anything else through unchanged."""
    text = (value or "").strip()
    if not text:
        return ""
    if text.lower() == "present":
        return "Present"
    try:
        # strptime parses a string into a datetime using a format code, then
        # strftime formats it back out. Raises ValueError if it doesn't match.
        return datetime.strptime(text, "%Y-%m").strftime("%b %Y")
    except ValueError:
        return text  # e.g. a bare year, or free-form text — leave it alone


def _date_range(start: str, end: str) -> str:
    return f"{_fmt_date(start)} – {_fmt_date(end)}"  # – is an en dash


def _org_label(employer: str, client: str | None) -> str:
    """Combine employer and client into one heading, e.g. "Acme (BigCo)".

    The client is rendered in parentheses, so a client name that already ends in
    its own parenthetical would nest them — "Acme (Dept. of X (DOX))". Strip that
    trailing group first. `r"..."` is a raw string (no backslash escaping), which
    is why regex patterns are always written that way in Python.
    """
    if not client:
        return employer
    cleaned = re.sub(r"\s*\([^()]*\)\s*$", "", client).strip()
    return f"{employer} ({cleaned or client})"


def _add_meta_line(doc, text: str) -> None:
    """Add a small italic line (used for role/date and project metadata)."""
    paragraph = doc.add_paragraph()
    # A "run" is a span of uniformly-formatted text inside a paragraph — you
    # need one to set italic, since formatting lives on the run, not the text.
    run = paragraph.add_run(text)
    run.italic = True


def render(
    profile: Profile,
    tailored: TailoredResume,
    out_path: Path,
    job_description: str = "",
) -> list[str]:
    """Write a tailored resume to `out_path`.

    Returns the list of `accomplishment_id`s the model referenced that don't
    exist in the profile — normally empty. The caller warns about them rather
    than dropping them silently.

    `job_description` only orders each role's skills line, putting the tags the
    posting mentions first. It defaults to empty so the renderer stays usable
    (and testable) without a posting; the line then keeps taxonomy order.
    """
    # Build id -> record lookups (a dict comprehension; like ToDictionary).
    accomplishments = {a.id: a for a in profile.accomplishments}
    projects = {p.id: p for p in profile.projects}

    # Group the selected bullets by the job they belong to. The key is a tuple
    # of the fields that identify one role; tuples are hashable, so they work
    # as dict keys (unlike a list).
    jobs: dict[tuple[str, str | None, str, str, str], list[str]] = {}
    project_bullets: dict[str, list[str]] = {}
    unmatched: list[str] = []

    # Every skill tag belonging to each role — drawn from *all* of that role's
    # accomplishments, not just the ones whose bullets made this cut. A skill you
    # used on the job is true whether or not its bullet won the selection.
    role_skills: dict[tuple[str, str | None, str, str, str], list[str]] = {}
    for acc in profile.accomplishments:
        role_key = (acc.employer, acc.client, acc.role, acc.start, acc.end)
        role_skills.setdefault(role_key, []).extend(acc.skills)

    job_tokens = skills_mod.job_tokens(job_description)
    taxonomy_order = profile.skill_taxonomy.all_tags()

    for bullet in tailored.bullets:
        acc = accomplishments.get(bullet.accomplishment_id)
        if acc is not None:
            key = (acc.employer, acc.client, acc.role, acc.start, acc.end)
            # setdefault returns the existing list or inserts a new one first.
            jobs.setdefault(key, []).append(bullet.text)
            continue

        if bullet.accomplishment_id in projects:
            project_bullets.setdefault(bullet.accomplishment_id, []).append(bullet.text)
            continue

        unmatched.append(bullet.accomplishment_id)

    doc = Document()
    c = profile.contact

    doc.add_heading(c.name, level=0)

    # A generator expression (like LINQ's .Where().Select()): keep only the
    # non-empty fields, then join them with " | ". Empty strings are falsy, so
    # `if p` drops any contact detail left blank in profile.yaml.
    headline = " | ".join(p for p in [c.title, c.tagline] if p)
    if headline:
        doc.add_paragraph(headline)

    contact = " | ".join(p for p in [c.email, c.phone, c.location] if p)
    if contact:
        doc.add_paragraph(contact)

    links = " | ".join(p for p in [c.linkedin, c.github] if p)
    if links:
        doc.add_paragraph(links)

    if tailored.summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(tailored.summary)

    # These are nested functions so they can read `doc`, `jobs`, and `projects`
    # from the enclosing scope (a closure — same idea as a C# lambda capturing
    # locals). Defining them here lets us emit the two sections in either order.
    def _render_experience() -> None:
        doc.add_heading("Experience", level=1)
        # Reverse-chronological: sort by start date, newest first. The lambda
        # picks the sort key out of each (key, bullets) pair; ISO-ish dates like
        # "2024-06" sort correctly as plain strings.
        for key, bullets in sorted(jobs.items(), key=lambda kv: kv[0][3], reverse=True):
            employer, client, role, start, end = key
            doc.add_heading(_org_label(employer, client), level=2)
            _add_meta_line(doc, f"{role}  |  {_date_range(start, end)}")
            names = skills_mod.skills_for_role(
                role_skills.get(key, []),
                profile.skill_display,
                taxonomy_order,
                job_tokens,
            )
            if names:
                _add_meta_line(doc, ", ".join(names))
            for text in bullets:
                doc.add_paragraph(text, style="List Bullet")

    def _render_projects() -> None:
        doc.add_heading("Projects", level=1)
        for pid, bullets in sorted(
            project_bullets.items(), key=lambda kv: projects[kv[0]].start, reverse=True
        ):
            project = projects[pid]
            doc.add_heading(project.name, level=2)
            meta = _date_range(project.start, project.end)
            if project.stack:
                # Same display map as the roles above, so a project's stack
                # doesn't render as raw slugs next to a job's "C#, .NET".
                stack = [skills_mod.display_name(t, profile.skill_display) for t in project.stack]
                meta = f"{', '.join(stack)}  |  {meta}"
            _add_meta_line(doc, meta)
            for text in bullets:
                doc.add_paragraph(text, style="List Bullet")

    # Order the two sections by recency so the whole document reads as one
    # unbroken reverse-chronological timeline. When projects are more recent
    # than the newest job — e.g. ongoing work covering a gap between roles —
    # they lead, so the reader sees current activity before the employment
    # history rather than after it.
    # `Callable[[], None]` is a function taking no args and returning nothing —
    # the type-hint equivalent of C#'s `Action`.
    sections: list[tuple[str, Callable[[], None]]] = []
    if jobs:
        sections.append((max(key[3] for key in jobs), _render_experience))
    if project_bullets:
        sections.append((max(projects[pid].start for pid in project_bullets), _render_projects))

    for _start, render_section in sorted(sections, key=lambda s: s[0], reverse=True):
        render_section()

    if profile.education:
        doc.add_heading("Education", level=1)
        for edu in profile.education:
            line = f"{edu.credential}, {edu.institution}"
            if edu.year:
                line = f"{line} ({edu.year})"
            doc.add_paragraph(line)

    # Make sure the output directory exists (parents=True is like -p; exist_ok
    # avoids an error if it already exists).
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return unmatched
